import io

import pytest

from app.models import Rule
from app.services import linter, registry
from app.services.llm import LLMError, LLMRouter, MockClient
from app.services.parsing import draft_rules_from_chunks, parse_document
from app.services.redaction import redact


def test_redaction_catches_common_pii():
    r = redact("Mail a.b@corp.com or call +91 98765 43210, PAN ABCDE1234F, Project Zephyr.", ["Zephyr"])
    for leaked in ("a.b@corp.com", "98765", "ABCDE1234F", "Zephyr"):
        assert leaked not in r.redacted
    assert len(r.spans) >= 4


class Exploding:
    model = "should-never-be-called"

    def generate(self, *a, **k):
        raise AssertionError("commercial API must not be called")


def test_restricted_never_leaves_system():
    router = LLMRouter(MockClient(), commercial=Exploding())
    router.run("explain", {"template": "x"}, classification="Restricted", external_confirmed=True)
    router.run("explain", {"template": "x"}, classification="Internal", external_confirmed=False)  # unconfirmed
    assert {c["provider"] for c in router.calls} == {"local"}


def test_external_payload_is_redacted_after_confirmation():
    seen = {}

    class Spy:
        model = "spy"

        def generate(self, task, payload, temperature=0.0):
            seen.update(payload)
            return {"explanation": "ok"}

    router = LLMRouter(MockClient(), commercial=Spy())
    router.run("explain", {"template": "Contact jo@x.com about 60%"}, classification="Public", external_confirmed=True)
    assert "jo@x.com" not in seen["template"] and router.calls[0]["provider"] == "commercial"


def test_commercial_failure_falls_back_to_local():
    class Down:
        model = "down"

        def generate(self, *a, **k):
            raise LLMError("offline")

    router = LLMRouter(MockClient(), commercial=Down())
    out = router.run("explain", {"template": "t"}, classification="Public", external_confirmed=True)
    assert out["explanation"] == "t" and router.calls[-1]["provider"] == "local"


def test_linter_finds_ambiguity_and_term_conflict(session, llm):
    registry.approve_rule(session, registry.create_rule(session, dict(
        rule_key="R-TERM-009", kind="required_term", text="Always include the phrase miracle in headlines.",
        pattern={"terms": ["miracle"]}), "m").id, "m")
    found = linter.lint_rules(session, llm)
    kinds = {(f["kind"], tuple(f["rules"])) for f in found}
    assert ("ambiguity", ("R-TONE-001",)) in kinds
    assert any(f["kind"] == "conflict" and set(f["rules"]) == {"R-TERM-001", "R-TERM-009"} for f in found)
    assert not any(f["kind"] == "ambiguity" and f["rules"] == ["R-TERM-001"] for f in found)


def test_parse_docx_and_pdf_and_draft_rules():
    import docx
    import fitz

    d = docx.Document()
    d.add_paragraph('Never use "world-class" in any headline copy.')
    d.add_paragraph('You must include "Terms apply" on offers and promotions.')
    buf = io.BytesIO()
    d.save(buf)
    parsed = parse_document("guide.docx", buf.getvalue())
    drafts = draft_rules_from_chunks(parsed.chunks)
    assert {(x["kind"], x["pattern"]["terms"][0]) for x in drafts} == {("banned_term", "world-class"),
                                                                       ("required_term", "Terms apply")}
    pdf = fitz.open()
    pdf.new_page().insert_text((72, 72), "Page one guideline paragraph about tone of voice.")
    parsed_pdf = parse_document("g.pdf", pdf.tobytes())
    assert "tone of voice" in parsed_pdf.text and parsed_pdf.chunks[0].ref.startswith("p.1")


def test_drafted_rules_are_inactive(session):
    r = registry.create_rule(session, dict(rule_key="D-1", kind="banned_term", text="x", pattern={"terms": ["x"]}),
                             "sw", drafted_by_software=True)
    assert r.status == "draft" and r.drafted_by_software
