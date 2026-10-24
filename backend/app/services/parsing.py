"""Document parsing for guideline and content uploads."""
import io
import re
from dataclasses import dataclass, field

from ..errors import DomainError


@dataclass
class Chunk:
    text: str
    ref: str
    page: int = 0
    index: int = 0


@dataclass
class ParsedDocument:
    filename: str
    text: str
    chunks: list = field(default_factory=list)


def _paragraphs(text: str):
    return [p.strip() for p in re.split(r"\n\s*\n|\n(?=[•\-\u2022]\s)", text) if len(p.strip()) >= 12]


def parse_document(filename: str, data: bytes) -> ParsedDocument:
    name = filename.lower()
    chunks, parts = [], []
    if name.endswith(".pdf"):
        import fitz  # PyMuPDF

        try:
            doc = fitz.open(stream=data, filetype="pdf")
        except Exception as e:
            raise DomainError(f"Could not read PDF: {e}")
        for pno, page in enumerate(doc, start=1):
            ptxt = page.get_text("text")
            parts.append(ptxt)
            for i, para in enumerate(_paragraphs(ptxt), start=1):
                chunks.append(Chunk(para, f"p.{pno} ¶{i}", pno, len(chunks)))
    elif name.endswith(".docx"):
        import docx

        try:
            d = docx.Document(io.BytesIO(data))
        except Exception as e:
            raise DomainError(f"Could not read DOCX: {e}")
        paras = [p.text.strip() for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            for row in t.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    paras.append(" | ".join(cells))
        parts = paras
        for i, para in enumerate(paras, start=1):
            if len(para) >= 12:
                chunks.append(Chunk(para, f"¶{i}", 0, len(chunks)))
    elif name.endswith((".txt", ".md")):
        txt = data.decode("utf-8", errors="replace")
        parts = [txt]
        for i, para in enumerate(_paragraphs(txt), start=1):
            chunks.append(Chunk(para, f"¶{i}", 0, len(chunks)))
    else:
        raise DomainError("Unsupported file type; use PDF, DOCX, TXT or MD")
    return ParsedDocument(filename, "\n\n".join(parts), chunks)


_BAN = re.compile(r"(?:never|do not|don'?t|avoid)\s+(?:use|say|write|claim)\s+[\u201c\"']([^\u201d\"']+)[\u201d\"']", re.I)
_REQ = re.compile(r"(?:always|must)\s+(?:include|use|add|say)\s+[\u201c\"']([^\u201d\"']+)[\u201d\"']", re.I)


def draft_rules_from_chunks(chunks: list) -> list:
    """Draft rules from quoted banned and required terms. Drafts remain inactive until approved."""
    drafts, n = [], 0
    for c in chunks:
        for rx, kind in ((_BAN, "banned_term"), (_REQ, "required_term")):
            for m in rx.finditer(c.text):
                n += 1
                drafts.append({
                    "rule_key": f"DRAFT-{n:03d}", "kind": kind, "text": c.text[:300], "severity": "medium",
                    "source_ref": c.ref, "pattern": {"terms": [m.group(1)]}, "drafted_by_software": True,
                })
    return drafts
